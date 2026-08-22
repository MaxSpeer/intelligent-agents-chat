"""Document parsing, lifecycle, hybrid retrieval, and project-isolation tests."""

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from intelligent_agents_chat.database import ChatRepository
from intelligent_agents_chat.documents import (
    BlobStore,
    Chunker,
    DocumentParser,
    DocumentStore,
    DocumentValidationError,
)
from intelligent_agents_chat.embeddings import EmbeddingError
from intelligent_agents_chat.rag import DocumentService, ProjectRAGRetriever
from intelligent_agents_chat.rag_evaluation import run_fixture_evaluation
from intelligent_agents_chat.retrieval import RetrievalQuery


class SemanticFakeEmbeddingGateway:
    model_name = "test-embedding"

    async def embed(self, texts):
        return [self._vector(text) for text in texts]

    @staticmethod
    def _vector(text: str) -> list[float]:
        normalized = text.casefold()
        if "orchid" in normalized or "flower" in normalized:
            return [1.0, 0.0, 0.0]
        if "finance" in normalized or "budget" in normalized:
            return [0.0, 1.0, 0.0]
        return [0.0, 0.0, 1.0]


class FlakyEmbeddingGateway(SemanticFakeEmbeddingGateway):
    def __init__(self) -> None:
        self.fail = True

    async def embed(self, texts):
        if self.fail:
            self.fail = False
            raise EmbeddingError("temporary embedding outage")
        return await super().embed(texts)


class DocumentRAGTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self) -> None:
        self.temporary_directory = TemporaryDirectory()
        root = Path(self.temporary_directory.name)
        self.database_path = root / "chats.sqlite3"
        self.repository = ChatRepository(self.database_path)
        self.repository.initialize()
        self.store = DocumentStore(self.database_path)
        self.store.initialize()
        self.blob_store = BlobStore(root / "documents")
        self.embedding_gateway = SemanticFakeEmbeddingGateway()
        self.service = DocumentService(
            self.store,
            self.blob_store,
            DocumentParser(),
            Chunker(chunk_chars=240, overlap_chars=40),
            self.embedding_gateway,
        )

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    async def test_upload_duplicate_replace_and_delete_have_deterministic_lifecycle(self) -> None:
        project = self.repository.create_project("Research")
        first = await self.service.upload(
            project_id=project.id,
            display_name="notes.txt",
            media_type="text/plain",
            data=b"The orchid launch code is amber.",
        )

        self.assertFalse(first.duplicate)
        self.assertEqual(first.document.status, "indexed")
        self.assertEqual(first.document.embedding_model, "test-embedding")
        self.assertTrue(self.blob_store.exists(first.document.storage_key))
        self.assertGreater(first.document.chunk_count, 0)

        duplicate = await self.service.upload(
            project_id=project.id,
            display_name="renamed.txt",
            media_type="text/plain",
            data=b"The orchid launch code is amber.",
        )
        self.assertTrue(duplicate.duplicate)
        self.assertEqual(duplicate.document.id, first.document.id)
        self.assertEqual(len(self.store.list_documents(project.id)), 1)

        replaced = await self.service.replace(
            first.document.id,
            project.id,
            display_name="updated.md",
            media_type="text/markdown",
            data=b"# Decision\nThe finance budget is cobalt.",
        )
        self.assertEqual(replaced.id, first.document.id)
        self.assertEqual(replaced.display_name, "updated.md")
        self.assertNotEqual(replaced.sha256, first.document.sha256)
        self.assertIn("section: Decision", self.store.list_chunks(replaced.id)[0].locator)

        self.assertTrue(await self.service.delete(replaced.id, project.id))
        self.assertIsNone(self.store.get_document(replaced.id))
        self.assertEqual(self.store.list_chunks(replaced.id), [])
        self.assertFalse(self.blob_store.exists(replaced.storage_key))

    async def test_hybrid_retrieval_never_crosses_project_boundaries_and_has_citations(
        self,
    ) -> None:
        public = self.repository.create_project("Public")
        private = self.repository.create_project("Private")
        public_upload = await self.service.upload(
            project_id=public.id,
            display_name="orchids.txt",
            media_type="text/plain",
            data=b"The orchid launch code is amber.",
        )
        await self.service.upload(
            project_id=private.id,
            display_name="secret.txt",
            media_type="text/plain",
            data=b"The orchid launch code is violet.",
        )

        retriever = ProjectRAGRetriever(self.store, self.embedding_gateway)
        results = await retriever.retrieve(
            RetrievalQuery(project_id=public.id, text="What is the flower launch code?")
        )

        self.assertGreaterEqual(len(results), 1)
        self.assertTrue(all(result.project_id == public.id for result in results))
        self.assertTrue(all("violet" not in result.text for result in results))
        self.assertEqual(results[0].source_kind, "project_document")
        self.assertEqual(results[0].title, "orchids.txt")
        self.assertEqual(results[0].locator, "document")
        self.assertEqual(results[0].source_conversation_id, None)
        self.assertIn(
            results[0].source_id,
            {chunk.id for chunk in self.store.list_chunks(public_upload.document.id)},
        )

    async def test_failed_embedding_is_visible_and_retryable(self) -> None:
        project = self.repository.create_project("Retries")
        flaky = FlakyEmbeddingGateway()
        service = DocumentService(
            self.store,
            self.blob_store,
            DocumentParser(),
            Chunker(),
            flaky,
        )

        with self.assertRaisesRegex(EmbeddingError, "temporary embedding outage"):
            await service.upload(
                project_id=project.id,
                display_name="retry.txt",
                media_type="text/plain",
                data=b"An orchid fact that should survive a temporary outage.",
            )

        failed = self.store.list_documents(project.id)
        self.assertEqual(len(failed), 1)
        self.assertEqual(failed[0].status, "failed")
        self.assertIn("temporary embedding outage", failed[0].error_message or "")
        self.assertTrue(self.blob_store.exists(failed[0].storage_key))

        indexed = await service.reindex(failed[0].id, project.id)
        self.assertEqual(indexed.status, "indexed")
        self.assertIsNone(indexed.error_message)
        self.assertGreater(indexed.chunk_count, 0)

    async def test_lexical_mode_works_without_an_embedding_endpoint(self) -> None:
        project = self.repository.create_project("Lexical")
        service = DocumentService(
            self.store,
            self.blob_store,
            DocumentParser(),
            Chunker(),
            None,
        )
        uploaded = await service.upload(
            project_id=project.id,
            display_name="sqlite.txt",
            media_type="text/plain",
            data=b"SQLite FTS5 provides deterministic lexical retrieval.",
        )

        self.assertIsNone(uploaded.document.embedding_model)
        results = await ProjectRAGRetriever(self.store, None).retrieve(
            RetrievalQuery(project_id=project.id, text="deterministic FTS5")
        )
        self.assertEqual(len(results), 1)
        self.assertIn("SQLite", results[0].text)

    async def test_xlsx_upload_is_searchable_with_a_worksheet_row_citation(self) -> None:
        project = self.repository.create_project("Spreadsheets")
        uploaded = await self.service.upload(
            project_id=project.id,
            display_name="launches.xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data=_xlsx_with_content(),
        )

        results = await ProjectRAGRetriever(self.store, self.embedding_gateway).retrieve(
            RetrievalQuery(project_id=project.id, text="What is the orchid launch code?")
        )

        self.assertEqual(uploaded.document.status, "indexed")
        self.assertEqual(results[0].title, "launches.xlsx")
        self.assertEqual(results[0].locator, 'sheet "Launches", row 2')
        self.assertIn("amber", results[0].text)

    async def test_checked_in_evaluation_set_measures_recall_before_reranking(self) -> None:
        result = await run_fixture_evaluation(k=2)

        self.assertEqual(result.case_count, 3)
        self.assertEqual(result.recall_at_k, 1.0)
        self.assertEqual(result.hit_rate_at_k, 1.0)
        self.assertGreater(result.mean_reciprocal_rank, 0.0)


class DocumentParserTests(unittest.TestCase):
    def test_chunking_is_stable_and_overlapping(self) -> None:
        chunker = Chunker(chunk_chars=200, overlap_chars=40)
        segments = DocumentParser().parse(
            ("alpha beta gamma delta " * 40).encode(),
            display_name="long.txt",
            media_type="text/plain",
        )

        first = chunker.chunk("document-id", segments)
        second = chunker.chunk("document-id", segments)

        self.assertGreater(len(first), 1)
        self.assertEqual([chunk.id for chunk in first], [chunk.id for chunk in second])
        self.assertEqual(
            [chunk.content_hash for chunk in first], [chunk.content_hash for chunk in second]
        )
        self.assertTrue(set(first[0].content.split()).intersection(first[1].content.split()))

    def test_markdown_sections_and_pdf_pages_preserve_locators(self) -> None:
        parser = DocumentParser()
        markdown = parser.parse(
            b"# Architecture\nUse SQLite.\n\n## Retrieval\nUse hybrid search.",
            display_name="design.md",
            media_type="text/markdown",
        )
        pdf = parser.parse(
            _pdf_with_text("Evidence on page one"),
            display_name="evidence.pdf",
            media_type="application/pdf",
        )

        self.assertEqual(
            [segment.locator for segment in markdown],
            ["section: Architecture", "section: Retrieval"],
        )
        self.assertEqual(pdf[0].locator, "page 1")
        self.assertIn("Evidence on page one", pdf[0].text)

    def test_docx_preserves_heading_paragraph_and_table_locators(self) -> None:
        parser = DocumentParser()
        segments = parser.parse(
            _docx_with_content(),
            display_name="research.docx",
            media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )

        self.assertEqual(segments[0].locator, "section: Research")
        self.assertIn("Orchid launch notes", segments[0].text)
        table_segment = next(segment for segment in segments if "table 1" in segment.locator)
        self.assertEqual(table_segment.locator, "section: Research, table 1, row 2")
        self.assertEqual(table_segment.text, "Project: Orchid | Code: amber")

    def test_csv_and_xlsx_preserve_structured_row_locators(self) -> None:
        parser = DocumentParser()
        csv_segments = parser.parse(
            b"Project;Launch code\nOrchid;amber\nFinance;cobalt\n",
            display_name="launches.csv",
            media_type="text/csv",
        )
        xlsx_segments = parser.parse(
            _xlsx_with_content(),
            display_name="launches.xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        self.assertEqual(csv_segments[0].locator, "row 2")
        self.assertEqual(csv_segments[0].text, "Project: Orchid | Launch code: amber")
        self.assertEqual(xlsx_segments[0].locator, 'sheet "Launches", row 2')
        self.assertEqual(xlsx_segments[0].text, "Project: Orchid | Launch code: amber")
        self.assertEqual(xlsx_segments[0].section, "Launches")

    def test_rejects_invalid_office_archives_and_mismatched_media_type(self) -> None:
        parser = DocumentParser()
        with self.assertRaisesRegex(DocumentValidationError, "not a valid DOCX"):
            parser.parse(
                b"not a zip archive",
                display_name="broken.docx",
                media_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            )
        with self.assertRaisesRegex(DocumentValidationError, "Unsupported media type for XLSX"):
            parser.parse(
                _xlsx_with_content(),
                display_name="launches.xlsx",
                media_type="application/pdf",
            )

    def test_rejects_unsupported_binary_and_empty_scanned_pdf(self) -> None:
        parser = DocumentParser()
        with self.assertRaisesRegex(DocumentValidationError, "Supported document types"):
            parser.parse(b"binary", display_name="image.png", media_type="image/png")

        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(DocumentValidationError, "no extractable text"):
            parser.parse(
                output.getvalue(),
                display_name="scan.pdf",
                media_type="application/pdf",
            )


def _pdf_with_text(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    page[NameObject("/Resources")] = resources
    stream = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream.set_data(f"BT /F1 12 Tf 72 200 Td ({escaped}) Tj ET".encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx_with_content() -> bytes:
    document = WordDocument()
    document.add_heading("Research", level=1)
    document.add_paragraph("Orchid launch notes")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Project"
    table.cell(0, 1).text = "Code"
    table.cell(1, 0).text = "Orchid"
    table.cell(1, 1).text = "amber"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_with_content() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Launches"
    worksheet.append(["Project", "Launch code"])
    worksheet.append(["Orchid", "amber"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


if __name__ == "__main__":
    unittest.main()
