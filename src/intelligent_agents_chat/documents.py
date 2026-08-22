"""Project-scoped document lifecycle, parsing, chunking, and SQLite persistence."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
import hashlib
from io import BytesIO, StringIO
import json
import logging
import math
from pathlib import Path
import re
import sqlite3
from typing import Callable, Iterable, Literal, Sequence
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader

from intelligent_agents_chat.database import PROJECT_ROOT, connect_database
from intelligent_agents_chat.logging_config import log_event


DEFAULT_DOCUMENT_ROOT = PROJECT_ROOT / ".data" / "documents"
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_DOCUMENTS_PER_PROJECT = 100
DEFAULT_CHUNK_CHARS = 1_200
DEFAULT_CHUNK_OVERLAP_CHARS = 180
MAX_QUERY_TERMS = 12
MAX_RETRIEVED_CHARS = 4_000
MAX_ARCHIVE_ENTRIES = 5_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_TABULAR_ROWS = 20_000
MAX_TABULAR_COLUMNS = 256
MAX_WORKBOOK_SHEETS = 50
SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx", ".csv", ".xlsx"}
GENERIC_MEDIA_TYPES = {"application/octet-stream", ""}
SUPPORTED_MEDIA_TYPES_BY_EXTENSION = {
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/x-markdown", "text/plain"},
    ".markdown": {"text/markdown", "text/x-markdown", "text/plain"},
    ".pdf": {"application/pdf", "application/x-pdf"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
        "application/zip",
    },
    ".csv": {
        "text/csv",
        "text/comma-separated-values",
        "application/csv",
        "application/vnd.ms-excel",
        "text/plain",
    },
    ".xlsx": {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
        "application/zip",
    },
}
SUPPORTED_MEDIA_TYPES = GENERIC_MEDIA_TYPES.union(
    *(media_types for media_types in SUPPORTED_MEDIA_TYPES_BY_EXTENSION.values())
)
DocumentStatus = Literal["processing", "indexed", "failed"]
logger = logging.getLogger(__name__)

_STOP_WORDS = {
    "aber",
    "auch",
    "das",
    "der",
    "die",
    "ein",
    "eine",
    "einer",
    "für",
    "ist",
    "mit",
    "oder",
    "the",
    "und",
    "von",
    "was",
    "wie",
    "wir",
    "zu",
}


class DocumentError(RuntimeError):
    """Base class for recoverable document lifecycle failures."""


class DocumentValidationError(DocumentError, ValueError):
    """An uploaded file violates a documented input constraint."""


@dataclass(frozen=True, slots=True)
class Document:
    id: str
    project_id: str
    display_name: str
    storage_key: str
    media_type: str
    extension: str
    sha256: str
    byte_size: int
    status: DocumentStatus
    error_message: str | None
    chunk_count: int
    embedding_model: str | None
    embedding_dimension: int | None
    created_at: datetime
    updated_at: datetime
    indexed_at: datetime | None


@dataclass(frozen=True, slots=True)
class ParsedSegment:
    text: str
    locator: str
    page_number: int | None = None
    section: str | None = None


@dataclass(frozen=True, slots=True)
class ChunkDraft:
    id: str
    ordinal: int
    content: str
    content_hash: str
    locator: str
    page_number: int | None
    section: str | None


@dataclass(frozen=True, slots=True)
class DocumentChunk:
    id: str
    document_id: str
    project_id: str
    ordinal: int
    content: str
    content_hash: str
    locator: str
    page_number: int | None
    section: str | None
    embedding: tuple[float, ...] | None
    embedding_model: str | None


class BlobStore:
    """Persist original uploads under generated project/document paths only."""

    def __init__(self, root: Path = DEFAULT_DOCUMENT_ROOT) -> None:
        self.root = root

    def storage_key(self, project_id: str, document_id: str, extension: str) -> str:
        _validate_storage_component(project_id)
        _validate_storage_component(document_id)
        normalized_extension = _normalized_extension(extension)
        return f"{project_id}/{document_id}/source{normalized_extension}"

    def write(self, storage_key: str, data: bytes) -> None:
        path = self._resolve(storage_key)
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(path.suffix + ".tmp")
        temporary.write_bytes(data)
        temporary.replace(path)

    def read(self, storage_key: str) -> bytes:
        return self._resolve(storage_key).read_bytes()

    def delete(self, storage_key: str) -> None:
        path = self._resolve(storage_key)
        path.unlink(missing_ok=True)
        document_directory = path.parent
        if document_directory.exists() and not any(document_directory.iterdir()):
            document_directory.rmdir()
        project_directory = document_directory.parent
        if project_directory.exists() and not any(project_directory.iterdir()):
            project_directory.rmdir()

    def exists(self, storage_key: str) -> bool:
        return self._resolve(storage_key).is_file()

    def _resolve(self, storage_key: str) -> Path:
        relative = Path(storage_key)
        if relative.is_absolute() or ".." in relative.parts:
            raise DocumentValidationError("Invalid generated document storage key")
        root = self.root.resolve()
        resolved = (root / relative).resolve()
        if resolved != root and root not in resolved.parents:
            raise DocumentValidationError("Document path escapes the configured storage root")
        return resolved


class DocumentParser:
    """Extract text and stable locators from supported document types."""

    def parse(self, data: bytes, *, display_name: str, media_type: str) -> list[ParsedSegment]:
        extension = validate_document_input(display_name, media_type, len(data))
        if extension == ".pdf":
            return self._parse_pdf(data)
        if extension == ".docx":
            return self._parse_docx(data)
        if extension == ".csv":
            return self._parse_csv(data)
        if extension == ".xlsx":
            return self._parse_xlsx(data)
        return self._parse_text(data, markdown=extension in {".md", ".markdown"})

    def _parse_text(self, data: bytes, *, markdown: bool) -> list[ParsedSegment]:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise DocumentValidationError("Text and Markdown files must use UTF-8") from error
        if "\x00" in text:
            raise DocumentValidationError("Text file contains binary null bytes")
        normalized = _normalize_document_text(text)
        if not normalized:
            raise DocumentValidationError("Document contains no extractable text")
        if not markdown:
            return [ParsedSegment(text=normalized, locator="document")]

        segments: list[ParsedSegment] = []
        current_section = "Document"
        current_lines: list[str] = []
        for line in normalized.splitlines():
            heading = re.match(r"^#{1,6}\s+(.+?)\s*$", line)
            if heading and current_lines:
                section_text = "\n".join(current_lines).strip()
                if section_text:
                    segments.append(
                        ParsedSegment(
                            text=section_text,
                            locator=f"section: {current_section}",
                            section=current_section,
                        )
                    )
                current_lines = []
            if heading:
                current_section = heading.group(1).strip()[:200]
            current_lines.append(line)
        section_text = "\n".join(current_lines).strip()
        if section_text:
            segments.append(
                ParsedSegment(
                    text=section_text,
                    locator=f"section: {current_section}",
                    section=current_section,
                )
            )
        return segments

    def _parse_pdf(self, data: bytes) -> list[ParsedSegment]:
        try:
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted and reader.decrypt("") == 0:
                raise DocumentValidationError("Password-protected PDFs are not supported")
            segments = []
            for page_index, page in enumerate(reader.pages, start=1):
                text = _normalize_document_text(page.extract_text() or "")
                if text:
                    segments.append(
                        ParsedSegment(
                            text=text,
                            locator=f"page {page_index}",
                            page_number=page_index,
                        )
                    )
        except DocumentValidationError:
            raise
        except Exception as error:
            raise DocumentValidationError("PDF could not be parsed safely") from error
        if not segments:
            raise DocumentValidationError(
                "PDF contains no extractable text; scanned PDFs require OCR"
            )
        return segments

    def _parse_docx(self, data: bytes) -> list[ParsedSegment]:
        _validate_office_archive(data, extension=".docx", required_member="word/document.xml")
        try:
            document = WordDocument(BytesIO(data))
            segments: list[ParsedSegment] = []
            paragraph_lines: list[str] = []
            current_section = "Document"
            table_index = 0
            tabular_rows = 0

            def flush_paragraphs() -> None:
                if not paragraph_lines:
                    return
                text = _normalize_document_text("\n\n".join(paragraph_lines))
                paragraph_lines.clear()
                if not text:
                    return
                segments.append(
                    ParsedSegment(
                        text=text,
                        locator=(
                            "document"
                            if current_section == "Document"
                            else f"section: {current_section}"
                        ),
                        section=None if current_section == "Document" else current_section,
                    )
                )

            for block in document.iter_inner_content():
                if isinstance(block, Paragraph):
                    text = _normalize_document_text(block.text)
                    if not text:
                        continue
                    style_name = getattr(getattr(block, "style", None), "name", "") or ""
                    if style_name.casefold().startswith("heading"):
                        flush_paragraphs()
                        current_section = text[:200]
                    paragraph_lines.append(text)
                    continue
                if not isinstance(block, Table):
                    continue

                flush_paragraphs()
                table_index += 1
                table_prefix = (
                    f"table {table_index}"
                    if current_section == "Document"
                    else f"section: {current_section}, table {table_index}"
                )
                table_segments, row_count = _tabular_segments(
                    (
                        (row_index, [cell.text for cell in row.cells])
                        for row_index, row in enumerate(block.rows, start=1)
                    ),
                    locator_for_row=lambda row_index, prefix=table_prefix: (
                        f"{prefix}, row {row_index}"
                    ),
                    section=None if current_section == "Document" else current_section,
                    max_data_rows=MAX_TABULAR_ROWS - tabular_rows,
                )
                tabular_rows += row_count
                segments.extend(table_segments)

            flush_paragraphs()
        except DocumentValidationError:
            raise
        except Exception as error:
            raise DocumentValidationError("DOCX could not be parsed safely") from error
        if not segments:
            raise DocumentValidationError("DOCX contains no extractable text")
        return segments

    def _parse_csv(self, data: bytes) -> list[ParsedSegment]:
        text = _decode_utf8(data, format_name="CSV")
        try:
            sample = text[:8192]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(StringIO(text, newline=""), dialect)

            def rows() -> Iterable[tuple[int, Sequence[object]]]:
                for row in reader:
                    yield reader.line_num, row

            segments, _ = _tabular_segments(
                rows(),
                locator_for_row=lambda row_number: f"row {row_number}",
                section="CSV",
                max_data_rows=MAX_TABULAR_ROWS,
            )
        except DocumentValidationError:
            raise
        except csv.Error as error:
            raise DocumentValidationError("CSV could not be parsed safely") from error
        if not segments:
            raise DocumentValidationError("CSV contains no extractable rows")
        return segments

    def _parse_xlsx(self, data: bytes) -> list[ParsedSegment]:
        _validate_office_archive(data, extension=".xlsx", required_member="xl/workbook.xml")
        workbook = None
        try:
            workbook = load_workbook(
                BytesIO(data),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
                raise DocumentValidationError(
                    f"XLSX exceeds the {MAX_WORKBOOK_SHEETS}-worksheet limit"
                )

            segments: list[ParsedSegment] = []
            total_rows = 0
            for worksheet in workbook.worksheets:
                if worksheet.max_column > MAX_TABULAR_COLUMNS:
                    raise DocumentValidationError(
                        f'Worksheet "{worksheet.title}" exceeds the '
                        f"{MAX_TABULAR_COLUMNS}-column limit"
                    )
                section = _normalize_cell_text(worksheet.title)[:200] or "Worksheet"
                sheet_segments, row_count = _tabular_segments(
                    enumerate(worksheet.iter_rows(values_only=True), start=1),
                    locator_for_row=lambda row_number, title=section: (
                        f'sheet "{title}", row {row_number}'
                    ),
                    section=section,
                    max_data_rows=MAX_TABULAR_ROWS - total_rows,
                )
                total_rows += row_count
                segments.extend(sheet_segments)
        except DocumentValidationError:
            raise
        except Exception as error:
            raise DocumentValidationError("XLSX could not be parsed safely") from error
        finally:
            if workbook is not None:
                workbook.close()
        if not segments:
            raise DocumentValidationError("XLSX contains no extractable cells")
        return segments


class Chunker:
    """Create deterministic, overlapping chunks without crossing source locators."""

    def __init__(
        self,
        *,
        chunk_chars: int = DEFAULT_CHUNK_CHARS,
        overlap_chars: int = DEFAULT_CHUNK_OVERLAP_CHARS,
    ) -> None:
        if chunk_chars < 200:
            raise ValueError("chunk_chars must be at least 200")
        if overlap_chars < 0 or overlap_chars >= chunk_chars:
            raise ValueError("overlap_chars must be non-negative and smaller than chunk_chars")
        self.chunk_chars = chunk_chars
        self.overlap_chars = overlap_chars

    def chunk(self, document_id: str, segments: Sequence[ParsedSegment]) -> list[ChunkDraft]:
        drafts: list[ChunkDraft] = []
        for segment in segments:
            pieces = _split_text(segment.text, self.chunk_chars, self.overlap_chars)
            for piece_index, piece in enumerate(pieces, start=1):
                content_hash = hashlib.sha256(piece.encode("utf-8")).hexdigest()
                ordinal = len(drafts)
                stable_seed = f"{document_id}\0{ordinal}\0{content_hash}".encode()
                chunk_id = hashlib.sha256(stable_seed).hexdigest()
                locator = segment.locator
                if len(pieces) > 1:
                    locator = f"{locator}, chunk {piece_index}/{len(pieces)}"
                drafts.append(
                    ChunkDraft(
                        id=chunk_id,
                        ordinal=ordinal,
                        content=piece,
                        content_hash=content_hash,
                        locator=locator,
                        page_number=segment.page_number,
                        section=segment.section,
                    )
                )
        if not drafts:
            raise DocumentValidationError("Document contains no indexable text chunks")
        return drafts


class DocumentStore:
    """Document metadata, chunks, lexical index, and optional vectors in SQLite."""

    def __init__(self, database_path: Path) -> None:
        self.database_path = database_path

    def initialize(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS documents (
                    id TEXT PRIMARY KEY,
                    project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                    display_name TEXT NOT NULL,
                    storage_key TEXT NOT NULL UNIQUE,
                    media_type TEXT NOT NULL,
                    extension TEXT NOT NULL,
                    sha256 TEXT NOT NULL,
                    byte_size INTEGER NOT NULL,
                    status TEXT NOT NULL,
                    error_message TEXT,
                    chunk_count INTEGER NOT NULL DEFAULT 0,
                    embedding_model TEXT,
                    embedding_dimension INTEGER,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    indexed_at TEXT,
                    UNIQUE(project_id, sha256)
                );

                CREATE INDEX IF NOT EXISTS documents_project_updated_idx
                    ON documents(project_id, updated_at DESC);
                CREATE INDEX IF NOT EXISTS documents_project_status_idx
                    ON documents(project_id, status);

                CREATE TABLE IF NOT EXISTS document_chunks (
                    row_id INTEGER PRIMARY KEY AUTOINCREMENT,
                    id TEXT NOT NULL UNIQUE,
                    document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                    ordinal INTEGER NOT NULL,
                    content TEXT NOT NULL,
                    content_hash TEXT NOT NULL,
                    locator TEXT NOT NULL,
                    page_number INTEGER,
                    section TEXT,
                    embedding_json TEXT,
                    embedding_model TEXT,
                    created_at TEXT NOT NULL,
                    UNIQUE(document_id, ordinal)
                );

                CREATE INDEX IF NOT EXISTS document_chunks_document_idx
                    ON document_chunks(document_id, ordinal);
                CREATE INDEX IF NOT EXISTS document_chunks_project_idx
                    ON document_chunks(project_id, document_id);

                CREATE VIRTUAL TABLE IF NOT EXISTS document_chunks_fts USING fts5(
                    content,
                    content='document_chunks',
                    content_rowid='row_id',
                    tokenize='unicode61 remove_diacritics 2'
                );

                CREATE TRIGGER IF NOT EXISTS document_chunks_after_insert
                AFTER INSERT ON document_chunks BEGIN
                    INSERT INTO document_chunks_fts(rowid, content)
                    VALUES (new.row_id, new.content);
                END;

                CREATE TRIGGER IF NOT EXISTS document_chunks_after_delete
                AFTER DELETE ON document_chunks BEGIN
                    INSERT INTO document_chunks_fts(document_chunks_fts, rowid, content)
                    VALUES ('delete', old.row_id, old.content);
                END;

                CREATE TRIGGER IF NOT EXISTS document_chunks_after_update
                AFTER UPDATE OF content ON document_chunks BEGIN
                    INSERT INTO document_chunks_fts(document_chunks_fts, rowid, content)
                    VALUES ('delete', old.row_id, old.content);
                    INSERT INTO document_chunks_fts(rowid, content)
                    VALUES (new.row_id, new.content);
                END;
                """
            )
        log_event(
            logger,
            logging.INFO,
            "documents.store.initialized",
            database_path=str(self.database_path),
        )

    def create_pending(
        self,
        *,
        project_id: str,
        display_name: str,
        storage_key: str,
        media_type: str,
        extension: str,
        sha256: str,
        byte_size: int,
        document_id: str | None = None,
    ) -> Document:
        now = _timestamp()
        created_id = document_id or str(uuid4())
        with self._connect() as connection:
            connection.execute(
                """
                INSERT INTO documents (
                    id, project_id, display_name, storage_key, media_type, extension,
                    sha256, byte_size, status, error_message, chunk_count,
                    embedding_model, embedding_dimension, created_at, updated_at, indexed_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'processing', NULL, 0, NULL, NULL, ?, ?, NULL)
                """,
                (
                    created_id,
                    project_id,
                    display_name,
                    storage_key,
                    media_type,
                    extension,
                    sha256,
                    byte_size,
                    now,
                    now,
                ),
            )
        document = self.get_document(created_id)
        if document is None:
            raise RuntimeError("Failed to read document after creating it")
        return document

    def find_duplicate(self, project_id: str, sha256: str) -> Document | None:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT * FROM documents WHERE project_id = ? AND sha256 = ?
                """,
                (project_id, sha256),
            ).fetchone()
        return _document_from_row(row) if row else None

    def get_document(self, document_id: str) -> Document | None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM documents WHERE id = ?",
                (document_id,),
            ).fetchone()
        return _document_from_row(row) if row else None

    def list_documents(self, project_id: str) -> list[Document]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT * FROM documents
                WHERE project_id = ?
                ORDER BY updated_at DESC, display_name COLLATE NOCASE, id
                """,
                (project_id,),
            ).fetchall()
        return [_document_from_row(row) for row in rows]

    def count_documents(self, project_id: str) -> int:
        with self._connect() as connection:
            return int(
                connection.execute(
                    "SELECT COUNT(*) FROM documents WHERE project_id = ?",
                    (project_id,),
                ).fetchone()[0]
            )

    def mark_processing(self, document_id: str) -> bool:
        with self._connect() as connection:
            cursor = connection.execute(
                """
                UPDATE documents
                SET status = 'processing', error_message = NULL, updated_at = ?
                WHERE id = ?
                """,
                (_timestamp(), document_id),
            )
        return cursor.rowcount == 1

    def update_original(
        self,
        document_id: str,
        *,
        display_name: str,
        storage_key: str,
        media_type: str,
        extension: str,
        sha256: str,
        byte_size: int,
    ) -> bool:
        with self._connect() as connection:
            cursor = connection.execute(
                """
                UPDATE documents
                SET display_name = ?, storage_key = ?, media_type = ?, extension = ?,
                    sha256 = ?, byte_size = ?, status = 'processing', error_message = NULL,
                    chunk_count = 0, embedding_model = NULL, embedding_dimension = NULL,
                    updated_at = ?, indexed_at = NULL
                WHERE id = ?
                """,
                (
                    display_name,
                    storage_key,
                    media_type,
                    extension,
                    sha256,
                    byte_size,
                    _timestamp(),
                    document_id,
                ),
            )
        return cursor.rowcount == 1

    def replace_chunks(
        self,
        document_id: str,
        chunks: Sequence[ChunkDraft],
        *,
        embeddings: Sequence[Sequence[float]] | None,
        embedding_model: str | None,
    ) -> Document:
        if embeddings is not None and len(embeddings) != len(chunks):
            raise ValueError("Embedding count must match chunk count")
        embedding_dimension = None
        if embeddings:
            dimensions = {len(vector) for vector in embeddings}
            if len(dimensions) != 1 or 0 in dimensions:
                raise ValueError("All embeddings must have one non-zero dimension")
            embedding_dimension = dimensions.pop()
        now = _timestamp()
        with self._connect() as connection:
            document = connection.execute(
                "SELECT project_id FROM documents WHERE id = ?",
                (document_id,),
            ).fetchone()
            if document is None:
                raise DocumentError("Document no longer exists")
            connection.execute("DELETE FROM document_chunks WHERE document_id = ?", (document_id,))
            connection.executemany(
                """
                INSERT INTO document_chunks (
                    id, document_id, project_id, ordinal, content, content_hash,
                    locator, page_number, section, embedding_json, embedding_model, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        chunk.id,
                        document_id,
                        document["project_id"],
                        chunk.ordinal,
                        chunk.content,
                        chunk.content_hash,
                        chunk.locator,
                        chunk.page_number,
                        chunk.section,
                        _serialize_vector(embeddings[index]) if embeddings is not None else None,
                        embedding_model if embeddings is not None else None,
                        now,
                    )
                    for index, chunk in enumerate(chunks)
                ],
            )
            connection.execute(
                """
                UPDATE documents
                SET status = 'indexed', error_message = NULL, chunk_count = ?,
                    embedding_model = ?, embedding_dimension = ?, updated_at = ?, indexed_at = ?
                WHERE id = ?
                """,
                (
                    len(chunks),
                    embedding_model if embeddings is not None else None,
                    embedding_dimension,
                    now,
                    now,
                    document_id,
                ),
            )
        result = self.get_document(document_id)
        if result is None:
            raise RuntimeError("Document disappeared after indexing")
        return result

    def mark_failed(self, document_id: str, error_message: str) -> bool:
        safe_message = " ".join(error_message.split())[:500] or "Unknown ingestion error"
        with self._connect() as connection:
            cursor = connection.execute(
                """
                UPDATE documents
                SET status = 'failed', error_message = ?, updated_at = ?
                WHERE id = ?
                """,
                (safe_message, _timestamp(), document_id),
            )
        return cursor.rowcount == 1

    def delete_document(self, document_id: str, project_id: str) -> Document | None:
        document = self.get_document(document_id)
        if document is None or document.project_id != project_id:
            return None
        with self._connect() as connection:
            connection.execute(
                "DELETE FROM documents WHERE id = ? AND project_id = ?",
                (document_id, project_id),
            )
        return document

    def list_chunks(self, document_id: str) -> list[DocumentChunk]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT id, document_id, project_id, ordinal, content, content_hash,
                       locator, page_number, section, embedding_json, embedding_model
                FROM document_chunks
                WHERE document_id = ?
                ORDER BY ordinal
                """,
                (document_id,),
            ).fetchall()
        return [_chunk_from_row(row) for row in rows]

    def lexical_search(
        self, project_id: str, text: str, limit: int
    ) -> list[tuple[DocumentChunk, float, str]]:
        fts_query = build_fts_query(text)
        if not fts_query or limit <= 0:
            return []
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT dc.id, dc.document_id, dc.project_id, dc.ordinal, dc.content,
                       dc.content_hash, dc.locator, dc.page_number, dc.section,
                       dc.embedding_json, dc.embedding_model, document.display_name,
                       bm25(document_chunks_fts) AS relevance
                FROM document_chunks_fts
                JOIN document_chunks AS dc ON dc.row_id = document_chunks_fts.rowid
                JOIN documents AS document ON document.id = dc.document_id
                WHERE document_chunks_fts MATCH ?
                  AND dc.project_id = ?
                  AND document.status = 'indexed'
                ORDER BY relevance ASC, dc.document_id, dc.ordinal
                LIMIT ?
                """,
                (fts_query, project_id, limit),
            ).fetchall()
        return [
            (_chunk_from_row(row), max(0.0, -float(row["relevance"])), row["display_name"])
            for row in rows
        ]

    def vector_rows(self, project_id: str) -> list[tuple[DocumentChunk, str]]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT dc.id, dc.document_id, dc.project_id, dc.ordinal, dc.content,
                       dc.content_hash, dc.locator, dc.page_number, dc.section,
                       dc.embedding_json, dc.embedding_model, document.display_name
                FROM document_chunks AS dc
                JOIN documents AS document ON document.id = dc.document_id
                WHERE dc.project_id = ?
                  AND document.status = 'indexed'
                  AND dc.embedding_json IS NOT NULL
                ORDER BY dc.document_id, dc.ordinal
                """,
                (project_id,),
            ).fetchall()
        return [(_chunk_from_row(row), row["display_name"]) for row in rows]

    def _connect(self) -> sqlite3.Connection:
        return connect_database(self.database_path)


def validate_document_input(display_name: str, media_type: str, byte_size: int) -> str:
    clean_name = Path(display_name).name.strip()
    if not clean_name:
        raise DocumentValidationError("Document filename cannot be empty")
    extension = _normalized_extension(Path(clean_name).suffix)
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentValidationError(
            "Supported document types are TXT, Markdown, PDF, DOCX, CSV, and XLSX"
        )
    normalized_media_type = media_type.split(";", 1)[0].strip().lower()
    allowed_media_types = GENERIC_MEDIA_TYPES | SUPPORTED_MEDIA_TYPES_BY_EXTENSION[extension]
    if normalized_media_type not in allowed_media_types:
        raise DocumentValidationError(
            f"Unsupported media type for {extension.removeprefix('.').upper()}: "
            f"{normalized_media_type or '(empty)'}"
        )
    if byte_size <= 0:
        raise DocumentValidationError("Document is empty")
    if byte_size > MAX_DOCUMENT_BYTES:
        raise DocumentValidationError(
            f"Document exceeds the {MAX_DOCUMENT_BYTES // (1024 * 1024)} MiB limit"
        )
    return extension


def build_fts_query(text: str) -> str:
    terms: list[str] = []
    seen: set[str] = set()
    for term in re.findall(r"[^\W_]+", text.casefold(), flags=re.UNICODE):
        if len(term) < 2 or term in _STOP_WORDS or term in seen:
            continue
        seen.add(term)
        terms.append(term)
        if len(terms) >= MAX_QUERY_TERMS:
            break
    return " OR ".join(f'"{term}"' for term in terms)


def cosine_similarity(left: Sequence[float], right: Sequence[float]) -> float:
    if len(left) != len(right) or not left:
        return 0.0
    numerator = sum(a * b for a, b in zip(left, right, strict=True))
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    if left_norm == 0 or right_norm == 0:
        return 0.0
    return numerator / (left_norm * right_norm)


def bounded_chunk_text(value: str) -> str:
    if len(value) <= MAX_RETRIEVED_CHARS:
        return value
    return value[: MAX_RETRIEVED_CHARS - 3].rstrip() + "..."


def _split_text(text: str, chunk_chars: int, overlap_chars: int) -> list[str]:
    normalized = text.strip()
    if len(normalized) <= chunk_chars:
        return [normalized] if normalized else []
    pieces: list[str] = []
    start = 0
    while start < len(normalized):
        maximum_end = min(len(normalized), start + chunk_chars)
        end = maximum_end
        if maximum_end < len(normalized):
            candidates = [
                normalized.rfind("\n\n", start + chunk_chars // 2, maximum_end),
                normalized.rfind("\n", start + chunk_chars // 2, maximum_end),
                normalized.rfind(" ", start + chunk_chars // 2, maximum_end),
            ]
            boundary = max(candidates)
            if boundary > start:
                end = boundary
        piece = normalized[start:end].strip()
        if piece:
            pieces.append(piece)
        if end >= len(normalized):
            break
        next_start = max(start + 1, end - overlap_chars)
        while next_start < end and not normalized[next_start].isspace():
            next_start += 1
        start = next_start
    return pieces


def _normalize_document_text(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in value.splitlines()]
    return "\n".join(lines).strip()


def _decode_utf8(data: bytes, *, format_name: str) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentValidationError(f"{format_name} files must use UTF-8") from error
    if "\x00" in text:
        raise DocumentValidationError(f"{format_name} file contains binary null bytes")
    return text


def _validate_office_archive(data: bytes, *, extension: str, required_member: str) -> None:
    format_name = extension.removeprefix(".").upper()
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            names = {member.filename for member in members}
            if required_member not in names:
                raise DocumentValidationError(f"File is not a valid {format_name} document")
            if len(members) > MAX_ARCHIVE_ENTRIES:
                raise DocumentValidationError(
                    f"{format_name} archive exceeds the {MAX_ARCHIVE_ENTRIES}-entry limit"
                )
            if any(member.flag_bits & 0x1 for member in members):
                raise DocumentValidationError(f"Encrypted {format_name} files are not supported")
            uncompressed_bytes = sum(member.file_size for member in members)
            if uncompressed_bytes > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise DocumentValidationError(
                    f"{format_name} expands beyond the "
                    f"{MAX_ARCHIVE_UNCOMPRESSED_BYTES // (1024 * 1024)} MiB safety limit"
                )
    except DocumentValidationError:
        raise
    except (BadZipFile, OSError) as error:
        raise DocumentValidationError(f"File is not a valid {format_name} document") from error


def _tabular_segments(
    rows: Iterable[tuple[int, Sequence[object]]],
    *,
    locator_for_row: Callable[[int], str],
    section: str | None,
    max_data_rows: int,
) -> tuple[list[ParsedSegment], int]:
    headers: list[str] | None = None
    header_row_number: int | None = None
    segments: list[ParsedSegment] = []
    data_row_count = 0

    for row_number, raw_values in rows:
        values = [_normalize_cell_text(value) for value in raw_values]
        if len(values) > MAX_TABULAR_COLUMNS:
            raise DocumentValidationError(
                f"Structured document exceeds the {MAX_TABULAR_COLUMNS}-column limit"
            )
        if not any(values):
            continue
        if headers is None:
            headers = _unique_headers(values)
            header_row_number = row_number
            continue
        if data_row_count >= max_data_rows:
            raise DocumentValidationError(
                f"Structured document exceeds the {MAX_TABULAR_ROWS}-data-row limit"
            )
        while len(headers) < len(values):
            headers.append(f"Column {len(headers) + 1}")
        fields = [
            f"{headers[index]}: {value}"
            for index, value in enumerate(values)
            if value
        ]
        if not fields:
            continue
        data_row_count += 1
        segments.append(
            ParsedSegment(
                text=" | ".join(fields),
                locator=locator_for_row(row_number),
                section=section,
            )
        )

    if headers is not None and not segments and header_row_number is not None:
        segments.append(
            ParsedSegment(
                text="Columns: " + ", ".join(headers),
                locator=locator_for_row(header_row_number),
                section=section,
            )
        )
    return segments, data_row_count


def _unique_headers(values: Sequence[str]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = value or f"Column {index}"
        occurrence = counts.get(base, 0) + 1
        counts[base] = occurrence
        headers.append(base if occurrence == 1 else f"{base} ({occurrence})")
    return headers


def _normalize_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def _normalized_extension(value: str) -> str:
    extension = value.casefold()
    if extension and not extension.startswith("."):
        extension = "." + extension
    if not re.fullmatch(r"\.[a-z0-9]+", extension):
        raise DocumentValidationError("Document extension is invalid")
    return extension


def _validate_storage_component(value: str) -> None:
    if not re.fullmatch(r"[A-Za-z0-9_-]+", value):
        raise DocumentValidationError("Generated storage identifier is invalid")


def _serialize_vector(values: Sequence[float]) -> str:
    vector = [float(value) for value in values]
    if not vector or not all(math.isfinite(value) for value in vector):
        raise ValueError("Embedding vectors must contain finite numbers")
    return json.dumps(vector, separators=(",", ":"))


def _deserialize_vector(value: str | None) -> tuple[float, ...] | None:
    if value is None:
        return None
    parsed = json.loads(value)
    if not isinstance(parsed, list) or not parsed:
        return None
    return tuple(float(item) for item in parsed)


def _document_from_row(row: sqlite3.Row) -> Document:
    return Document(
        id=row["id"],
        project_id=row["project_id"],
        display_name=row["display_name"],
        storage_key=row["storage_key"],
        media_type=row["media_type"],
        extension=row["extension"],
        sha256=row["sha256"],
        byte_size=int(row["byte_size"]),
        status=row["status"],
        error_message=row["error_message"],
        chunk_count=int(row["chunk_count"]),
        embedding_model=row["embedding_model"],
        embedding_dimension=row["embedding_dimension"],
        created_at=datetime.fromisoformat(row["created_at"]),
        updated_at=datetime.fromisoformat(row["updated_at"]),
        indexed_at=datetime.fromisoformat(row["indexed_at"]) if row["indexed_at"] else None,
    )


def _chunk_from_row(row: sqlite3.Row) -> DocumentChunk:
    return DocumentChunk(
        id=row["id"],
        document_id=row["document_id"],
        project_id=row["project_id"],
        ordinal=int(row["ordinal"]),
        content=row["content"],
        content_hash=row["content_hash"],
        locator=row["locator"],
        page_number=row["page_number"],
        section=row["section"],
        embedding=_deserialize_vector(row["embedding_json"]),
        embedding_model=row["embedding_model"],
    )


def _timestamp() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")
